"""End-to-end tests for the file endpoints.

These drive the real HTTP surface with real files: the PDF, DOCX and XLSX
fixtures are *generated* rather than committed as binary blobs, so the suite
proves the parsers work on genuine files without carrying opaque test data
nobody can review.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.file import File, FileStatus
from app.services import storage

API = "/api/v1"
EMAIL = "ada@example.com"
PASSWORD = "a-sufficiently-long-password"


# --- fixtures -------------------------------------------------------------


async def sign_up(client: AsyncClient, email: str = EMAIL) -> None:
    response = await client.post(
        f"{API}/auth/signup", json={"email": email, "password": PASSWORD}
    )
    assert response.status_code == 201


async def upload(
    client: AsyncClient,
    content: bytes,
    filename: str,
    content_type: str = "application/octet-stream",
) -> Any:
    return await client.post(
        f"{API}/files",
        files={"file": (filename, content, content_type)},
    )


def make_pdf(text: str = "Hello from a real PDF", pages: int = 1) -> bytes:
    """Build a genuinely valid PDF whose pages carry an extractable text layer.

    Assembled byte by byte rather than with a generator library, for two
    reasons: it adds no test-only dependency, and it makes the thing being
    tested visible. A PDF's "text layer" is the `BT ... Tj ET` operators in a
    page's content stream, together with a font resource -- a page without
    both is what a scanned document looks like, and yields nothing.

    The cross-reference table records the byte offset of every object, so the
    offsets are computed while writing rather than guessed.
    """
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        """Append an object and return its 1-based number."""
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    # Reserve the Pages object number so pages can reference their parent.
    pages_number = len(objects) + 1 + (2 * pages)

    page_numbers: list[int] = []
    for index in range(pages):
        stream = f"BT /F1 24 Tf 72 700 Td ({text} {index + 1}) Tj ET".encode("latin-1")
        contents = add(
            b"<< /Length "
            + str(len(stream)).encode()
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )
        page_numbers.append(
            add(
                b"<< /Type /Page /Parent "
                + str(pages_number).encode()
                + b" 0 R /MediaBox [0 0 612 792] /Contents "
                + str(contents).encode()
                + b" 0 R /Resources << /Font << /F1 "
                + str(font).encode()
                + b" 0 R >> >> >>"
            )
        )

    kids = b" ".join(f"{number} 0 R".encode() for number in page_numbers)
    add(
        b"<< /Type /Pages /Kids ["
        + kids
        + b"] /Count "
        + str(pages).encode()
        + b" >>"
    )
    catalog = add(b"<< /Type /Catalog /Pages " + str(pages_number).encode() + b" 0 R >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(number).encode() + b" 0 obj\n" + body + b"\nendobj\n"

    xref_at = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root "
        + str(catalog).encode()
        + b" 0 R >>\nstartxref\n"
        + str(xref_at).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """Build a real .docx in memory."""
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                word_table.cell(row_index, cell_index).text = value

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_xlsx(sheets: dict[str, list[list[Any]]]) -> bytes:
    """Build a real .xlsx in memory."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def make_png(width: int = 4, height: int = 4) -> bytes:
    """Build a real PNG in memory."""
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (width, height), (10, 20, 30)).save(out, format="PNG")
    return out.getvalue()


# --- authentication -------------------------------------------------------


class TestRequiresAuthentication:
    @pytest.mark.parametrize(
        ("method", "path"),
        [
            ("POST", "/files"),
            ("GET", "/files"),
            ("GET", "/files/11111111-1111-1111-1111-111111111111"),
            ("GET", "/files/11111111-1111-1111-1111-111111111111/download"),
            ("DELETE", "/files/11111111-1111-1111-1111-111111111111"),
        ],
    )
    async def test_every_endpoint_needs_a_session(
        self, client: AsyncClient, method: str, path: str
    ) -> None:
        response = await client.request(method, f"{API}{path}")

        assert response.status_code == 401


# --- text formats ---------------------------------------------------------


class TestTextUploads:
    async def test_stores_a_text_file_and_extracts_it(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)

        response = await upload(client, b"The quick brown fox.", "notes.txt")

        assert response.status_code == 201, response.text
        body = response.json()
        assert body["filename"] == "notes.txt"
        assert body["content_type"] == "text/plain"
        assert body["status"] == "ready"
        assert body["extracted_text"] == "The quick brown fox."
        assert body["has_text"] is True

    async def test_records_the_real_size(self, client: AsyncClient) -> None:
        await sign_up(client)
        content = b"x" * 5000

        body = (await upload(client, content, "notes.txt")).json()

        assert body["size_bytes"] == 5000

    async def test_extracts_a_csv_as_readable_rows(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)
        csv_bytes = b"name,role\nAda,engineer\nGrace,admiral\n"

        body = (await upload(client, csv_bytes, "people.csv")).json()

        assert body["content_type"] == "text/csv"
        # Rendered with a separator that survives values containing commas.
        assert "Ada | engineer" in body["extracted_text"]
        assert "Grace | admiral" in body["extracted_text"]

    async def test_markdown_is_stored_as_text(self, client: AsyncClient) -> None:
        await sign_up(client)

        body = (await upload(client, b"# Title\n\nBody.", "readme.md")).json()

        assert body["content_type"] == "text/plain"
        assert "# Title" in body["extracted_text"]

    async def test_a_utf8_document_survives_intact(
        self, client: AsyncClient
    ) -> None:
        """Multi-byte characters must not be mangled by chunked reads."""
        await sign_up(client)
        text = "Grüße — naïve café 日本語 🎉"

        body = (await upload(client, text.encode("utf-8"), "notes.txt")).json()

        assert body["extracted_text"] == text


# --- office formats -------------------------------------------------------


class TestOfficeUploads:
    async def test_extracts_docx_paragraphs(self, client: AsyncClient) -> None:
        await sign_up(client)
        docx_bytes = make_docx(["First paragraph.", "Second paragraph."])

        body = (await upload(client, docx_bytes, "report.docx")).json()

        assert body["status"] == "ready"
        assert "First paragraph." in body["extracted_text"]
        assert "Second paragraph." in body["extracted_text"]

    async def test_extracts_docx_tables(self, client: AsyncClient) -> None:
        """In real documents the tables carry the facts. Reading only
        paragraphs looks fine on a memo and drops the substance of a report."""
        await sign_up(client)
        docx_bytes = make_docx(
            ["Summary."], table=[["Region", "Revenue"], ["EMEA", "421000"]]
        )

        body = (await upload(client, docx_bytes, "report.docx")).json()

        assert "Region | Revenue" in body["extracted_text"]
        assert "EMEA | 421000" in body["extracted_text"]

    async def test_extracts_xlsx_cells_across_sheets(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)
        xlsx_bytes = make_xlsx(
            {
                "Sales": [["Region", "Revenue"], ["EMEA", 421000]],
                "Costs": [["Item", "Amount"], ["Cloud", 12000]],
            }
        )

        body = (await upload(client, xlsx_bytes, "figures.xlsx")).json()

        assert body["status"] == "ready"
        assert "EMEA | 421000" in body["extracted_text"]
        assert "Cloud | 12000" in body["extracted_text"]
        assert body["page_count"] == 2

    async def test_a_docx_renamed_as_xlsx_is_rejected(
        self, client: AsyncClient
    ) -> None:
        """Both are ZIPs, so the magic bytes cannot separate them -- but the
        parser can, and a mislabelled file must not be stored as the wrong
        type."""
        await sign_up(client)
        docx_bytes = make_docx(["Content."])

        response = await upload(client, docx_bytes, "actually-a-doc.xlsx")

        # Accepted by signature (it is a ZIP) but fails to parse as a workbook.
        assert response.status_code == 201
        assert response.json()["status"] == "failed"

    async def test_a_corrupt_office_file_is_kept_and_marked_failed(
        self, client: AsyncClient
    ) -> None:
        """A file we cannot parse is still the user's file. Deleting it out
        from under them would be worse than telling them the truth."""
        await sign_up(client)
        broken = io.BytesIO()
        with zipfile.ZipFile(broken, "w") as archive:
            archive.writestr("not-a-document.txt", "nonsense")

        response = await upload(client, broken.getvalue(), "broken.docx")

        assert response.status_code == 201
        body = response.json()
        assert body["status"] == "failed"
        assert body["error"]
        # SECURITY: the message is written for the user, not copied from the
        # exception -- parser errors quote file paths and library internals.
        assert "Traceback" not in body["error"]
        assert ".venv" not in body["error"]


# --- pdf ------------------------------------------------------------------


class TestPdfUploads:
    async def test_extracts_pdf_text(self, client: AsyncClient) -> None:
        await sign_up(client)
        pdf_bytes = make_pdf("Quarterly results")

        body = (await upload(client, pdf_bytes, "report.pdf")).json()

        assert body["status"] == "ready", body
        assert "Quarterly results" in body["extracted_text"]

    async def test_reports_the_page_count(self, client: AsyncClient) -> None:
        await sign_up(client)
        pdf_bytes = make_pdf("Page", pages=3)

        body = (await upload(client, pdf_bytes, "report.pdf")).json()

        assert body["page_count"] == 3

    async def test_a_pdf_with_no_text_layer_says_so(
        self, client: AsyncClient
    ) -> None:
        """A scan is a container of images. That is not a failure, and saying
        'no text' is more useful than an empty string that looks like a bug."""
        from pypdf import PdfWriter

        await sign_up(client)
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        out = io.BytesIO()
        writer.write(out)

        body = (await upload(client, out.getvalue(), "scan.pdf")).json()

        assert body["status"] == "extracted_empty"
        assert body["has_text"] is False
        assert "scan" in body["error"].lower()

    async def test_a_corrupt_pdf_is_kept_and_marked_failed(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)

        response = await upload(client, b"%PDF-1.7\nnot really a pdf", "broken.pdf")

        assert response.status_code == 201
        assert response.json()["status"] == "failed"


# --- images ---------------------------------------------------------------


class TestImageUploads:
    async def test_stores_an_image_without_extracting_text(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)

        body = (await upload(client, make_png(), "chart.png")).json()

        assert body["content_type"] == "image/png"
        assert body["status"] == "unsupported"
        assert body["has_text"] is False
        # Honest about *why* -- not a silent empty result.
        assert "vision" in body["error"].lower()

    async def test_a_fake_image_is_rejected(self, client: AsyncClient) -> None:
        """Magic bytes prove how a file starts, not what the rest of it is.
        Pillow parsing the header is what establishes it is really an image."""
        await sign_up(client)
        fake = b"\x89PNG\r\n\x1a\n" + b"\x00" * 200

        response = await upload(client, fake, "fake.png")

        assert response.status_code == 201
        assert response.json()["status"] == "failed"


# --- rejection ------------------------------------------------------------


class TestRejection:
    async def test_rejects_an_executable(self, client: AsyncClient) -> None:
        await sign_up(client)

        response = await upload(client, b"MZ\x90\x00\x03\x00", "setup.exe")

        assert response.status_code == 400

    async def test_rejects_an_html_file_disguised_as_an_image(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)

        response = await upload(
            client, b"<!DOCTYPE html><script>alert(1)</script>", "logo.png"
        )

        assert response.status_code == 400

    async def test_rejects_an_empty_file(self, client: AsyncClient) -> None:
        await sign_up(client)

        response = await upload(client, b"", "empty.txt")

        assert response.status_code == 400

    async def test_rejects_a_file_over_the_size_limit(
        self, client: AsyncClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The limit is counted while the bytes arrive, not read from the
        Content-Length header the client supplies."""
        await sign_up(client)
        from app.core.config import get_settings

        settings = get_settings()
        monkeypatch.setattr(settings, "max_upload_bytes", 1024, raising=False)

        response = await upload(client, b"x" * 5000, "big.txt")

        assert response.status_code == 400
        assert "larger than" in response.json()["detail"]

    async def test_nothing_is_stored_when_an_upload_is_rejected(
        self, client: AsyncClient, db: AsyncSession
    ) -> None:
        """A rejected upload must leave neither a row nor bytes behind."""
        await sign_up(client)

        await upload(client, b"MZ\x90\x00\x03\x00", "setup.exe")

        assert (await db.execute(select(File))).scalars().all() == []
        from tests.conftest import TEST_UPLOAD_DIR

        assert list(TEST_UPLOAD_DIR.iterdir()) == []

    async def test_is_rate_limited(self, client: AsyncClient) -> None:
        await sign_up(client)

        statuses = []
        for index in range(25):
            response = await upload(client, f"file {index}".encode(), "notes.txt")
            statuses.append(response.status_code)

        assert 429 in statuses, f"never rate limited: {set(statuses)}"


# --- filenames ------------------------------------------------------------


class TestFilenameHandling:
    async def test_a_traversing_filename_is_reduced_to_its_last_component(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)

        body = (await upload(client, b"content", "../../../etc/passwd.txt")).json()

        assert body["filename"] == "passwd.txt"

    async def test_the_stored_path_never_uses_the_users_filename(
        self, client: AsyncClient, db: AsyncSession
    ) -> None:
        """The real defence: the user's name is a label, and the path is built
        only from a key this application generated."""
        await sign_up(client)

        await upload(client, b"content", "../../escape.txt")

        stored = (await db.execute(select(File))).scalar_one()
        assert stored.storage_key.isalnum()
        assert len(stored.storage_key) == 32
        # The bytes live directly under the upload root, nowhere else.
        from tests.conftest import TEST_UPLOAD_DIR

        path = storage.path_for(stored.storage_key)
        assert path.parent == TEST_UPLOAD_DIR
        assert path.exists()

    async def test_a_filename_with_no_name_gets_one(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)

        body = (await upload(client, b"some text content", "...")).json()

        assert body["filename"] == "unnamed"


class TestStorageKeyValidation:
    @pytest.mark.parametrize(
        "key",
        [
            "../../../etc/passwd",
            "..\\..\\windows\\system32",
            "not-hex-at-all-but-32-characters",
            "abc",
            "",
        ],
    )
    def test_a_path_cannot_be_built_from_an_untrusted_key(self, key: str) -> None:
        """`path_for` is the single place a path is constructed, so validating
        here means no caller can turn a bad key into a filesystem reach."""
        with pytest.raises(ValueError, match="untrusted key"):
            storage.path_for(key)


# --- listing, reading, deleting -------------------------------------------


class TestListing:
    async def test_lists_newest_first(self, client: AsyncClient) -> None:
        await sign_up(client)
        await upload(client, b"first", "one.txt")
        await upload(client, b"second", "two.txt")

        page = (await client.get(f"{API}/files")).json()

        assert [item["filename"] for item in page["items"]] == ["two.txt", "one.txt"]

    async def test_paginates(self, client: AsyncClient) -> None:
        await sign_up(client)
        for index in range(3):
            await upload(client, f"file {index}".encode(), f"file{index}.txt")

        page = (await client.get(f"{API}/files", params={"limit": 2})).json()

        assert len(page["items"]) == 2
        assert page["has_more"] is True

    async def test_the_list_omits_the_full_text(self, client: AsyncClient) -> None:
        """The full text of a 500-page PDF has no business in a sidebar
        payload."""
        await sign_up(client)
        await upload(client, b"a" * 5000, "big.txt")

        page = (await client.get(f"{API}/files")).json()

        assert "extracted_text" not in page["items"][0]
        assert page["items"][0]["has_text"] is True

    async def test_the_storage_key_is_never_exposed(
        self, client: AsyncClient
    ) -> None:
        """It is the value that maps to a path on the server. The client has
        no use for it, so there is no reason to hand it out."""
        await sign_up(client)
        await upload(client, b"content", "notes.txt")

        listed = (await client.get(f"{API}/files")).text
        detail = (await client.get(f"{API}/files")).json()["items"][0]
        one = (await client.get(f"{API}/files/{detail['id']}")).text

        assert "storage_key" not in listed
        assert "storage_key" not in one


class TestReading:
    async def test_returns_the_extracted_text(self, client: AsyncClient) -> None:
        await sign_up(client)
        uploaded = (await upload(client, b"full document text", "notes.txt")).json()

        body = (await client.get(f"{API}/files/{uploaded['id']}")).json()

        assert body["extracted_text"] == "full document text"

    async def test_preview_truncates(self, client: AsyncClient) -> None:
        await sign_up(client)
        uploaded = (await upload(client, b"x" * 2000, "notes.txt")).json()

        body = (
            await client.get(f"{API}/files/{uploaded['id']}", params={"preview": True})
        ).json()

        assert len(body["extracted_text"]) == 400

    async def test_downloads_as_an_attachment(self, client: AsyncClient) -> None:
        """SECURITY: serving user content inline from the API's own origin is
        how stored XSS happens. These headers are the control."""
        await sign_up(client)
        uploaded = (await upload(client, b"the original bytes", "notes.txt")).json()

        response = await client.get(f"{API}/files/{uploaded['id']}/download")

        assert response.status_code == 200
        assert response.content == b"the original bytes"
        assert response.headers["content-type"] == "application/octet-stream"
        assert "attachment" in response.headers["content-disposition"]
        assert response.headers["x-content-type-options"] == "nosniff"

    async def test_a_missing_file_is_a_404(self, client: AsyncClient) -> None:
        await sign_up(client)

        response = await client.get(
            f"{API}/files/11111111-1111-1111-1111-111111111111"
        )

        assert response.status_code == 404


class TestDeleting:
    async def test_removes_the_row_and_the_bytes(
        self, client: AsyncClient, db: AsyncSession
    ) -> None:
        await sign_up(client)
        uploaded = (await upload(client, b"content", "notes.txt")).json()
        stored = (await db.execute(select(File))).scalar_one()
        path = storage.path_for(stored.storage_key)
        assert path.exists()

        response = await client.delete(f"{API}/files/{uploaded['id']}")

        assert response.status_code == 204
        # Deleting the row alone would leave the user's content on disk after
        # they asked for it to be gone.
        assert not path.exists()

    async def test_deleting_twice_is_a_404(self, client: AsyncClient) -> None:
        await sign_up(client)
        uploaded = (await upload(client, b"content", "notes.txt")).json()

        await client.delete(f"{API}/files/{uploaded['id']}")
        again = await client.delete(f"{API}/files/{uploaded['id']}")

        assert again.status_code == 404


# --- user isolation -------------------------------------------------------


class TestUserIsolation:
    async def _upload_as_ada_then_switch(self, client: AsyncClient) -> str:
        await sign_up(client, "ada@example.com")
        uploaded = (await upload(client, b"ada's private notes", "ada.txt")).json()
        await client.post(f"{API}/auth/logout")
        await sign_up(client, "mallory@example.com")
        return uploaded["id"]

    async def test_another_users_file_is_not_listed(
        self, client: AsyncClient
    ) -> None:
        await self._upload_as_ada_then_switch(client)

        page = (await client.get(f"{API}/files")).json()

        assert page["items"] == []

    async def test_another_users_file_cannot_be_read(
        self, client: AsyncClient
    ) -> None:
        file_id = await self._upload_as_ada_then_switch(client)

        response = await client.get(f"{API}/files/{file_id}")

        assert response.status_code == 404

    async def test_another_users_file_cannot_be_downloaded(
        self, client: AsyncClient
    ) -> None:
        file_id = await self._upload_as_ada_then_switch(client)

        response = await client.get(f"{API}/files/{file_id}/download")

        assert response.status_code == 404

    async def test_another_users_file_cannot_be_deleted(
        self, client: AsyncClient, db: AsyncSession
    ) -> None:
        file_id = await self._upload_as_ada_then_switch(client)

        response = await client.delete(f"{API}/files/{file_id}")

        assert response.status_code == 404
        assert len((await db.execute(select(File))).scalars().all()) == 1


# --- extraction limits ----------------------------------------------------


class TestExtractionLimits:
    async def test_a_long_document_is_truncated_and_says_so(
        self, client: AsyncClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """DOCX and XLSX are ZIP archives, so a small upload can hold an
        enormous amount of text. Extraction is bounded, and the user is told
        rather than silently handed a partial document."""
        await sign_up(client)
        from app.core.config import get_settings

        monkeypatch.setattr(get_settings(), "max_extracted_characters", 100)

        body = (await upload(client, b"y" * 5000, "long.txt")).json()

        assert len(body["extracted_text"]) == 100
        assert "longer than" in body["error"]

    async def test_a_pdf_beyond_the_page_cap_is_bounded(
        self, client: AsyncClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A hostile file with a huge page count is a CPU denial of service."""
        await sign_up(client)
        from app.core.config import get_settings

        monkeypatch.setattr(get_settings(), "max_pdf_pages", 2)
        pdf_bytes = make_pdf("Page", pages=5)

        body = (await upload(client, pdf_bytes, "long.pdf")).json()

        assert body["page_count"] == 5
        assert "longer than" in (body["error"] or "")


# --- account deletion -----------------------------------------------------


class TestAccountDeletion:
    async def test_deleting_a_user_removes_their_file_rows(
        self, client: AsyncClient, db: AsyncSession
    ) -> None:
        """The foreign key cascade covers the metadata only. The bytes on disk
        are a separate store with no foreign keys, so they are knowingly left
        behind -- recorded as a gap on the model and in ROADMAP.md rather than
        pretended away here."""
        from app.models.user import User

        await sign_up(client)
        await upload(client, b"content", "notes.txt")

        user = (await db.execute(select(User))).scalar_one()
        await db.delete(user)
        await db.commit()

        assert (await db.execute(select(File))).scalars().all() == []


# --- duplicate uploads ------------------------------------------------------


class TestDuplicateUploads:
    """Uploading bytes the user already has reuses the stored file.

    Without this, each copy is extracted and embedded again, and every search
    returns the same passage once per copy.
    """

    async def test_the_same_bytes_twice_reuse_the_first_file(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)
        first = await upload(client, b"quarterly report", "report.txt")
        second = await upload(client, b"quarterly report", "report.txt")

        assert first.status_code == 201
        assert first.json()["reused"] is False
        # 200, not 201: nothing was created.
        assert second.status_code == 200
        assert second.json()["reused"] is True
        assert second.json()["id"] == first.json()["id"]

    async def test_stores_one_row_and_one_file_on_disk(
        self, client: AsyncClient, db: AsyncSession
    ) -> None:
        from tests.conftest import TEST_UPLOAD_DIR

        await sign_up(client)
        for _ in range(3):
            await upload(client, b"quarterly report", "report.txt")

        assert len((await db.execute(select(File))).scalars().all()) == 1
        assert len(list(TEST_UPLOAD_DIR.iterdir())) == 1

    async def test_records_the_content_hash(
        self, client: AsyncClient, db: AsyncSession
    ) -> None:
        import hashlib

        await sign_up(client)
        await upload(client, b"quarterly report", "report.txt")

        [stored] = (await db.execute(select(File))).scalars().all()
        assert stored.content_sha256 == hashlib.sha256(b"quarterly report").hexdigest()

    async def test_a_new_name_does_not_make_a_new_file(
        self, client: AsyncClient
    ) -> None:
        """Identity is the bytes. The existing file keeps its original name."""
        await sign_up(client)
        await upload(client, b"quarterly report", "report.txt")
        second = await upload(client, b"quarterly report", "copy of report.txt")

        assert second.json()["reused"] is True
        assert second.json()["filename"] == "report.txt"

    async def test_different_bytes_are_different_files(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)
        first = await upload(client, b"version one", "report.txt")
        second = await upload(client, b"version two", "report.txt")

        assert second.status_code == 201
        assert second.json()["id"] != first.json()["id"]

    async def test_a_mismatched_extension_is_still_rejected(
        self, client: AsyncClient
    ) -> None:
        """Type checks run before the duplicate check, so renaming a stored
        file to a wrong extension is refused, not matched to its twin."""
        await sign_up(client)
        await upload(client, make_pdf(), "paper.pdf")

        renamed = await upload(client, make_pdf(), "paper.txt")

        assert renamed.status_code == 400

    async def test_another_users_identical_file_is_never_reused(
        self, client: AsyncClient
    ) -> None:
        """SECURITY: reusing across accounts would hand Mallory Ada's file id,
        and tell her that Ada holds this exact document."""
        await sign_up(client, "ada@example.com")
        ada = (await upload(client, b"shared memo", "memo.txt")).json()
        await client.post(f"{API}/auth/logout")
        await sign_up(client, "mallory@example.com")

        mallory = await upload(client, b"shared memo", "memo.txt")

        assert mallory.status_code == 201
        assert mallory.json()["reused"] is False
        assert mallory.json()["id"] != ada["id"]

    async def test_after_deleting_the_same_bytes_upload_fresh(
        self, client: AsyncClient
    ) -> None:
        await sign_up(client)
        first = (await upload(client, b"quarterly report", "report.txt")).json()
        await client.delete(f"{API}/files/{first['id']}")

        again = await upload(client, b"quarterly report", "report.txt")

        assert again.status_code == 201
        assert again.json()["id"] != first["id"]
